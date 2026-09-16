"""
Cross-check independen buat kelengkapan pipeline DWH -> Bronze -> Silver
Tier 1 -> Silver Tier 2 - TIDAK pakai kolom Bronze/Silver yang sudah
tercatat di Coretan, murni dihitung ulang dari 3 sumber file mapping
teknis terpisah:

  Step 1 (DWH -> Bronze):        dwh_to_flat.csv
  Step 2 (Bronze -> Silver 1):   Source ke Silver 1.csv
  Step 3 (Silver 1 -> Silver 2): Silver 1 ke Silver 2/*.xlsx (7 file per domain)

Catatan penting:
- Satu (Table, Column) DWH bisa punya LEBIH DARI SATU mapping Bronze
  (misal kolom yang sama dipetakan ke sumber Health & Life sekaligus) -
  jadi hasil akhirnya bisa >1 baris per kolom DWH (fan-out), bukan
  selalu 1:1. Step 3 dibatasi 1 hasil per key (berhenti di match
  pertama, sesuai urutan file & sheet).
- Pencarian bisa mulai dari stage MANAPUN (bukan cuma dari DWH), jadi
  `build_full_lineage_table()` juga menyertakan entri Bronze/Silver1
  yang "yatim" (orphan) - ada di file Step 2/3 tapi tidak ketarik dari
  DWH manapun di Step 1. Baris orphan itu kolom stage sebelumnya diisi
  "-".
"""

from __future__ import annotations

import csv
import os

import openpyxl
import pandas as pd

DWH_TO_FLAT_PATH = "dwh_to_flat.csv"
SRC_TO_SLV1_PATH = "Source ke Silver 1.csv"
SLV1_TO_SLV2_DIR = "Silver 1 ke Silver 2"

SLV1_TO_SLV2_FILES = [
    ("Application", "AMFS Data Model - SLV1 to SLV2 - Application.xlsx"),
    ("Claim", "AMFS Data Model - SLV1 to SLV2 - Claim.xlsx"),
    ("FinAct", "AMFS Data Model - SLV1 to SLV2 - FinAct.xlsx"),
    ("Party", "AMFS Data Model - SLV1 to SLV2 - Party.xlsx"),
    ("Policy", "AMFS Data Model - SLV1 to SLV2 - Policy.xlsx"),
    ("Producer", "AMFS Data Model - SLV1 to SLV2 - Producer.xlsx"),
    ("Product", "AMFS Data Model - SLV1 to SLV2 - Product.xlsx"),
]

LINEAGE_COLUMNS = [
    "DWH Table",
    "DWH Column",
    "Bronze Table",
    "Bronze Column",
    "Silver1 Table",
    "Silver1 Column",
    "Silver2 Table",
    "Silver2 Column",
    "Domain",
]


def load_bronze_to_silver1(path: str = SRC_TO_SLV1_PATH) -> tuple[dict, dict]:
    """Return (lookup, original_case):
      lookup: (BRONZE_TABLE, BRONZE_COL) uppercase -> list of (Silver1 Table, Silver1 Column)
      original_case: (BRONZE_TABLE, BRONZE_COL) uppercase -> (Bronze Table, Bronze Column) versi asli

    File py2-baris-header: baris 1 = section (Source/Bronze Layer/Silver 1),
    baris 2 = nama field. Kolom G/H (index 6/7) = Bronze Table/Field Name,
    kolom N (index 13) = Proposed Table Name (prioritas), M (index 12) =
    Table Name mentah (fallback), O (index 14) = Physical Field NM."""
    lookup: dict = {}
    original_case: dict = {}
    with open(path, encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    for r in rows[2:]:
        if len(r) < 15:
            continue
        bronze_table_raw = (r[6] or "").strip()
        bronze_col_raw = (r[7] or "").strip()
        if not bronze_table_raw or not bronze_col_raw:
            continue
        key = (bronze_table_raw.upper(), bronze_col_raw.upper())
        original_case.setdefault(key, (bronze_table_raw, bronze_col_raw))

        slv1_table = (r[13] or "").strip() or (r[12] or "").strip()
        slv1_col = (r[14] or "").strip()
        if not slv1_table:
            continue
        lookup.setdefault(key, [])
        pair = (slv1_table, slv1_col)
        if pair not in lookup[key]:
            lookup[key].append(pair)
    return lookup, original_case


def load_silver1_to_silver2(base_dir: str = SLV1_TO_SLV2_DIR) -> tuple[dict, dict]:
    """Return (lookup, original_case):
      lookup: (SILVER1_TABLE, SILVER1_COL) uppercase -> (Silver2 Table, Silver2 Column, Domain)
      original_case: (SILVER1_TABLE, SILVER1_COL) uppercase -> (Silver1 Table, Silver1 Column) versi asli

    Cuma 1 hasil per key di `lookup` - berhenti di match pertama sesuai
    urutan file (SLV1_TO_SLV2_FILES) & urutan sheet dalam file itu (sheet
    'Home' dilewati). Header selalu di baris 4, data mulai baris 5; kolom
    B/C (index 1/2) = Silver1 Table/Field Name, E/F (index 4/5) = Silver2
    Table Name/Physical Field Name."""
    lookup: dict = {}
    original_case: dict = {}
    for domain, filename in SLV1_TO_SLV2_FILES:
        path = os.path.join(base_dir, filename)
        if not os.path.exists(path):
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            if sheet_name == "Home":
                continue
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=5, values_only=True):
                if len(row) < 6:
                    continue
                b, c, e, f = row[1], row[2], row[4], row[5]
                if not b or not c:
                    continue
                key = (str(b).strip().upper(), str(c).strip().upper())
                original_case.setdefault(key, (str(b).strip(), str(c).strip()))
                if key not in lookup:
                    lookup[key] = (
                        str(e).strip() if e else "",
                        str(f).strip() if f else "",
                        domain,
                    )
        wb.close()
    return lookup, original_case


def build_full_lineage_table(
    dwh_path: str = DWH_TO_FLAT_PATH,
    src_to_slv1_path: str = SRC_TO_SLV1_PATH,
    slv1_to_slv2_dir: str = SLV1_TO_SLV2_DIR,
) -> pd.DataFrame:
    """Materialisasi SEMUA jalur lineage dari ke-3 sumber file, termasuk
    entri "yatim" (orphan) yang tidak ketarik dari DWH manapun - supaya
    bisa dicari/filter mulai dari stage APAPUN (DWH, Bronze, Silver1,
    atau Silver2), bukan cuma dari DWH.

    3 golongan baris yang dihasilkan:
      1. Rooted dari DWH (dwh_to_flat.csv) - DWH Table/Column terisi.
      2. Orphan Bronze - ada di Source ke Silver 1.csv tapi Bronze-nya
         tidak pernah jadi Source_Table/Column di dwh_to_flat.csv. DWH
         Table/Column = "-".
      3. Orphan Silver1 - ada di file Silver1->Silver2 tapi Silver1-nya
         tidak pernah jadi hasil Step 2 (baik dari jalur DWH maupun
         orphan Bronze di atas). DWH & Bronze Table/Column = "-".
    """
    bronze_to_silver1, bronze_case = load_bronze_to_silver1(src_to_slv1_path)
    silver1_to_silver2, silver1_case = load_silver1_to_silver2(slv1_to_slv2_dir)

    rows = []
    seen_bronze_keys = set()
    seen_silver1_keys = set()

    # 1. Rooted dari DWH
    with open(dwh_path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            dwh_table = (row.get("Target_Table") or "").strip()
            dwh_col = (row.get("Target_Column") or "").strip()
            b_table = (row.get("Source_Table") or "").strip()
            b_col = (row.get("Source_Column") or "").strip()
            if not dwh_table or not dwh_col or not b_table:
                continue

            b_key = (b_table.upper(), b_col.upper())
            seen_bronze_keys.add(b_key)
            silver1_list = bronze_to_silver1.get(b_key, [])
            if not silver1_list:
                rows.append(
                    {
                        "DWH Table": dwh_table,
                        "DWH Column": dwh_col,
                        "Bronze Table": b_table,
                        "Bronze Column": b_col,
                        "Silver1 Table": "-",
                        "Silver1 Column": "-",
                        "Silver2 Table": "-",
                        "Silver2 Column": "-",
                        "Domain": "-",
                    }
                )
                continue

            for s1_table, s1_col in silver1_list:
                s1_key = (s1_table.upper(), s1_col.upper())
                seen_silver1_keys.add(s1_key)
                s2 = silver1_to_silver2.get(s1_key)
                s2_table, s2_col, domain = s2 if s2 else ("-", "-", "-")
                rows.append(
                    {
                        "DWH Table": dwh_table,
                        "DWH Column": dwh_col,
                        "Bronze Table": b_table,
                        "Bronze Column": b_col,
                        "Silver1 Table": s1_table,
                        "Silver1 Column": s1_col,
                        "Silver2 Table": s2_table or "-",
                        "Silver2 Column": s2_col or "-",
                        "Domain": domain or "-",
                    }
                )

    # 2. Orphan Bronze: ada di Step 2 tapi Bronze-nya tidak pernah ketarik dari DWH
    for b_key, silver1_list in bronze_to_silver1.items():
        if b_key in seen_bronze_keys:
            continue
        b_table_orig, b_col_orig = bronze_case[b_key]
        for s1_table, s1_col in silver1_list:
            s1_key = (s1_table.upper(), s1_col.upper())
            seen_silver1_keys.add(s1_key)
            s2 = silver1_to_silver2.get(s1_key)
            s2_table, s2_col, domain = s2 if s2 else ("-", "-", "-")
            rows.append(
                {
                    "DWH Table": "-",
                    "DWH Column": "-",
                    "Bronze Table": b_table_orig,
                    "Bronze Column": b_col_orig,
                    "Silver1 Table": s1_table,
                    "Silver1 Column": s1_col,
                    "Silver2 Table": s2_table or "-",
                    "Silver2 Column": s2_col or "-",
                    "Domain": domain or "-",
                }
            )

    # 3. Orphan Silver1: ada di Step 3 tapi tidak pernah ketarik dari Step 2 (baik dari DWH maupun orphan Bronze)
    for s1_key, (s2_table, s2_col, domain) in silver1_to_silver2.items():
        if s1_key in seen_silver1_keys:
            continue
        s1_table_orig, s1_col_orig = silver1_case[s1_key]
        rows.append(
            {
                "DWH Table": "-",
                "DWH Column": "-",
                "Bronze Table": "-",
                "Bronze Column": "-",
                "Silver1 Table": s1_table_orig,
                "Silver1 Column": s1_col_orig,
                "Silver2 Table": s2_table or "-",
                "Silver2 Column": s2_col or "-",
                "Domain": domain or "-",
            }
        )

    return pd.DataFrame(rows, columns=LINEAGE_COLUMNS)
